"""Builds the Applied AI (22AI741) case-study report in the DSCE format, as .docx (+ PDF via LibreOffice).

Run:  python docs/build_report_aai.py
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(DOCS))

from studio import safety  # noqa: E402
from studio.consistency import RULES  # noqa: E402
from studio.game_bible import default_bible  # noqa: E402
from studio.vocab import ABILITY_POOL, ARC_TEMPLATES, PERSONALITIES, SITUATIONS  # noqa: E402

A = DOCS / "report_assets"
S = DOCS / "screenshots" / "crops"
OUT_DOCX = DOCS / "Case_Study_65_AAI_Report.docx"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

TITLE = "GENAI FOR VIDEO GAME CHARACTER DESIGN"
SHORT_TITLE = "GenAI for Video Game Character Design"
STUDENT, USN = "SURAJ T", "1DS23AI058"
GUIDE, GUIDE_DESIG = "Dr Reshma S", "Associate Professor"
HOD = "Dr. Vindhya P Malagi"
PRINCIPAL = "Dr. B G Prasad"
COURSE = "Applied AI:22AI741"
YEAR = "2026-27"

EV = json.loads((A / "eval.json").read_text())
ROWS, GUARD = EV["rows"], EV["guard"]
HI = ROWS[-1]
SAVING = HI["large_model_cost_per_char_usd"] / HI["cost_per_char_usd"]
BIBLE = default_bible()

BLUE = RGBColor(0x1F, 0x3B, 0x8C)
ORANGE = RGBColor(0xC5, 0x5A, 0x11)
RULE_COLOR = "843C0C"

# ----------------------------------------------------------------------------- low-level helpers
FONT = "Times New Roman"


def set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def rich(p, text, size=12, bold=False, italic=False, color=None, name=FONT):
    """Adds text to paragraph p; **bold** spans are bolded."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = p.add_run(part)
            set_font(r, size, bold or (i % 2 == 1), italic, color, name)
    return p


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.5, keep=False, indent=None, first=None):
    pf = p.paragraph_format
    p.alignment = align
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if keep:
        pf.keep_with_next = True
    if indent is not None:
        pf.left_indent = Inches(indent)
    if first is not None:
        pf.first_line_indent = Inches(first)
    return p


def para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, italic=False, color=None, before=0,
         after=6, line=1.5, keep=False):
    p = doc.add_paragraph()
    fmt(p, align, before, after, line, keep)
    rich(p, text, size, bold, italic, color)
    return p


def center(doc, text, size=12, bold=False, italic=False, color=None, before=0, after=0, line=1.0):
    return para(doc, text, WD_ALIGN_PARAGRAPH.CENTER, size, bold, italic, color, before, after, line)


def bullet(doc, text, size=12):
    p = doc.add_paragraph(style="List Bullet")
    fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2, 1.5)
    rich(p, text, size)
    return p


def numbered(doc, n, text, size=12):
    p = doc.add_paragraph()
    fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2, 1.5, indent=0.4, first=-0.3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
    rich(p, f"{n}.\t{text}", size)
    return p


def picture(doc, path, width):
    p = doc.add_paragraph()
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 1.0, keep=True)
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    fmt(p, after=0, line=1.0)


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(doc, header, rows, widths, size=11, header_fill="D9E2F3", align_cols=None, bold_first=False):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            al = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or (align_cols and align_cols[j] == "c")) else WD_ALIGN_PARAGRAPH.LEFT
            fmt(p, al, 2, 2, 1.1)
            rich(p, str(val), size, bold=(i == 0) or (bold_first and j == 0))
            if i == 0 and header_fill:
                shade(c, header_fill)
    fix_layout(t, widths)
    # repeat header row on page breaks
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def merge(a, b):
    """Merge cells and drop the empty paragraphs python-docx carries over."""
    m = a.merge(b)
    for extra in m.paragraphs[1:]:
        if not extra.text.strip():
            extra._element.getparent().remove(extra._element)
    return m


def fix_layout(t, widths):
    """Fixed layout with explicit grid widths (LibreOffice/Word otherwise auto-size) and unsplittable rows."""
    t.autofit = False
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trpr.insert(0, cs)
        for j, c in enumerate(row.cells):
            if j < len(widths):
                c.width = Inches(widths[j])


def borderless(t):
    tblpr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def para_border(p, side, val="thinThickSmallGap", sz=18, color=RULE_COLOR, space=1):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), val)
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pbdr.append(e)


def field(p, instr, size=10, color=None, name="Calibri"):
    for kind in ("begin", None, "separate", "1", "end"):
        r = p.add_run()
        set_font(r, size, color=color, name=name)
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif kind is None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r._r.append(it)
        else:
            r.text = kind


def sect_insert(sectPr, el, before_tags):
    for tag in before_tags:
        found = sectPr.find(qn(tag))
        if found is not None:
            found.addprevious(el)
            return
    sectPr.append(el)


def page_border(section):
    pb = OxmlElement("w:pgBorders")
    pb.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "thickThinSmallGap")
        e.set(qn("w:sz"), "24")
        e.set(qn("w:space"), "24")
        e.set(qn("w:color"), "auto")
        pb.append(e)
    sect_insert(section._sectPr, pb, ["w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign",
                                      "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
                                      "w:docGrid"])


def restart_numbering(section, start=1):
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), str(start))
    sect_insert(section._sectPr, pg, ["w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                                      "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid"])


def style_headings(doc):
    for name, size, al in (("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", 12, None),
                           ("Heading 3", 12, None)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.italic = FONT, Pt(size), True, False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in list(rf.attrib):
            del rf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(10 if name != "Heading 1" else 4), Pt(4), 1.5
        pf.keep_with_next = True
        if al is not None:
            pf.alignment = al


# ----------------------------------------------------------------------------- document state
class Report:
    def __init__(self, pages):
        self.pages = pages  # key -> main-section page number (from previous pass)
        self.figs, self.tabs, self.toc = [], [], []
        self.doc = Document()
        d = self.doc
        ns = d.styles["Normal"]
        ns.font.name, ns.font.size = FONT, Pt(12)
        rpr = ns.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        style_headings(d)
        s = d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin, s.right_margin = Inches(1.15), Inches(1.0)
        s.header_distance = s.footer_distance = Inches(0.45)
        page_border(s)

    # headings ---------------------------------------------------------------
    def chapter(self, n, title, first=False):
        d = self.doc
        if not first:
            page_break(d)
        para(d, f"CHAPTER {n}", WD_ALIGN_PARAGRAPH.LEFT, 12, bold=True, after=2, keep=True)
        h = d.add_heading(title.upper(), level=1)
        fmt(h, WD_ALIGN_PARAGRAPH.CENTER, 4, 10, 1.5, keep=True)
        self.toc.append((0, f"{n}  {title}", f"CHAPTER {n}"))

    def h2(self, num, title, toc=True):
        h = self.doc.add_heading(f"{num}  {title}", level=2)
        if toc:
            self.toc.append((1, f"{num} {title}", f"{num}  {title}"))

    def h3(self, num, title):
        self.doc.add_heading(f"{num}  {title}", level=3)
        self.toc.append((2, f"{num} {title}", None))

    def fig(self, path, width, num, caption):
        picture(self.doc, path, width)
        center(self.doc, f"Fig {num} {caption}", 11, before=4, after=10)
        self.figs.append((f"Fig {num}", caption, f"Fig {num} {caption}"))

    def tab(self, num, caption):
        p = center(self.doc, f"Table {num}: {caption}", 11, before=6, after=4)
        p.paragraph_format.keep_with_next = True
        self.tabs.append((num, caption, f"Table {num}: {caption}"))

    def pg(self, key):
        v = self.pages.get(key)
        return f"{v:02d}" if v else "--"


# ----------------------------------------------------------------------------- front matter
def college_block(d, blue=True, logo_top=False):
    if logo_top:
        picture(d, A / "logo_dsce.png", 0.95)
    c = BLUE if blue else None
    center(d, "Dayananda Sagar College of Engineering", 15, bold=True, color=c, before=6)
    center(d, "(An Autonomous Institute affiliated to Visvesvaraya Technological University (VTU), Belagavi, Approved by "
              "AICTE and UGC, Accredited by NAAC with 'A' grade)", 9)
    center(d, "Shavige Malleshwara Hills, Kumaraswamy Layout, Bengaluru - 560 111. India", 9, after=10)
    center(d, "DEPARTMENT OF", 12, bold=True, color=c)
    center(d, "ARTIFICIAL INTELLIGENCE & MACHINE LEARNING", 12, bold=True, color=c, after=8)


def front_matter(R):
    d = R.doc
    # ---- cover
    center(d, "DAYANANDA SAGAR COLLEGE OF ENGINEERING", 14, bold=True, before=0)
    center(d, "(An Autonomous Institute affiliated to Visvesvaraya Technological University (VTU), Belagavi,", 8)
    center(d, "Approved by AICTE and UGC, Accredited by NAAC with 'A' grade & ISO 9001 – 2015 Certified Institution)", 8)
    center(d, "Shavige Malleshwara Hills, Kumaraswamy Layout, Bengaluru-560 111, India", 9, after=14)
    center(d, "DEPARTMENT OF ARTIFICIAL INTELLIGENCE & MACHINE LEARNING", 12, bold=True, after=10)
    picture(d, A / "logo_vtu.png", 1.15)
    center(d, COURSE, 12, bold=True, before=14, after=12)
    center(d, "CASE STUDY REPORT ON", 12, bold=True, after=12)
    center(d, TITLE, 14, bold=True, after=20)
    center(d, "Submitted by", 12, bold=True, italic=True, after=8)
    center(d, f"{STUDENT} – {USN}", 12, after=26)
    center(d, "Under the Guidance of", 12, bold=True, italic=True, after=2)
    for line in (GUIDE, GUIDE_DESIG, "Department of Artificial Intelligence & Machine Learning", "DSCE, Bengaluru"):
        center(d, line, 11)
    d.add_paragraph().paragraph_format.space_after = Pt(6)
    picture(d, A / "logo_dsce.png", 0.95)
    center(d, "VISVESVARAYA TECHNOLOGICAL UNIVERSITY", 12, bold=True, before=18)
    center(d, "JNANASANGAMA, BELAGAVI-590018, KARNATAKA, INDIA", 12, bold=True)
    center(d, YEAR, 12, bold=True)

    # ---- certificate
    page_break(d)
    college_block(d)
    picture(d, A / "logo_dsce.png", 0.95)
    center(d, "CERTIFICATE", 15, bold=True, color=ORANGE, before=16, after=16)
    para(d, f"This is to certify that the Project work entitled **“{TITLE}”** carried out by **{STUDENT} ({USN})**, "
            "is the Bonafide work of students of **Dayananda Sagar College of Engineering** submitted in partial "
            "fulfillment for the award of **Bachelor of Engineering** in **Artificial Intelligence & Machine Learning, "
            f"affiliated under Visvesvaraya Technological University, Belagavi**, during the year **{YEAR}**. It is "
            "certified that all corrections/suggestions indicated for internal assessment have been incorporated in the "
            "report and deposited in the department library. This Case Study report has been approved as it satisfies "
            "the academic requirements in respect of AAT component of Applied AI.", size=11.5, line=1.6)
    for _ in range(3):
        d.add_paragraph()
    t = d.add_table(rows=3, cols=2)
    borderless(t)
    for j, (name, des) in enumerate([(GUIDE, GUIDE_DESIG), (HOD, "Professor & HOD")]):
        for i, (txt, b) in enumerate([("--------------------------------", False), (name, True), (des, False)]):
            p = t.cell(i, j).paragraphs[0]
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 1.2)
            rich(p, txt, 11, bold=b)

    # ---- declaration
    page_break(d)
    picture(d, A / "logo_dsce.png", 0.95)
    college_block(d)
    center(d, "DECLARATION", 14, bold=True, before=20, after=16)
    para(d, f"I, **{STUDENT} ({USN})** student of the 7th semester B.E., Artificial Intelligence & Machine Learning, "
            "Dayananda Sagar College of Engineering, Bangalore, declare that this case study work entitled "
            f"**“{TITLE}”** has been carried by us and submitted in the partial fulfilment of the Applied AI course "
            "requirements for the award of Degree in **Bachelor of Engineering** in **Artificial Intelligence & Machine "
            f"Learning** by **Visvesvaraya Technological University, Belagavi** during the year **{YEAR}.**",
         size=11.5, line=1.6)
    for _ in range(4):
        d.add_paragraph()
    para(d, "**Signature of the Student**", WD_ALIGN_PARAGRAPH.RIGHT, 11, after=2, line=1.2)
    para(d, f"**{STUDENT} ({USN})**", WD_ALIGN_PARAGRAPH.RIGHT, 11, line=1.2)

    # ---- vision & mission
    page_break(d)
    center(d, "INSTITUTIONAL & DEPARTMENTAL VISION AND MISSION", 14, bold=True, before=6, after=14, line=1.2)
    blocks = [
        ("Institute Vision", ["To impart quality technical education with a focus on Research and Innovation emphasizing "
                              "on Development of Sustainable and Inclusive Technology for the benefit of society."], False),
        ("Institute Mission", ["To provide an environment that enhances creativity and Innovation in pursuit of excellence.",
                               "To nurture team work in order to transform individuals as responsible leaders and entrepreneurs.",
                               "To train the students to the changing technical scenario and make them to understand the "
                               "importance of Sustainable and Inclusive technologies."], True),
        ("Department Vision", ["To provide progressive education and flourish the students' ingenuity to be successful "
                               "professionals impacting the society for a smarter and ethical world."], False),
        ("Department Mission", ["To adopt an engaging teaching–learning process with emphasis on problem-solving, "
                                "programming, and intelligent systems development skills.",
                                "To promote additional skill development through enhanced and experiential learning.",
                                "To collaborate with industries and professional bodies to make students industry-ready.",
                                "To encourage innovation through multi-disciplinary research and development activities.",
                                "To nurture human values and ethical principles in students, empowering them to contribute "
                                "meaningfully to society as responsible professionals."], True),
    ]
    for head, items, is_list in blocks:
        para(d, head, WD_ALIGN_PARAGRAPH.LEFT, 13, bold=True, before=8, after=4, keep=True)
        for it in items:
            (bullet(d, it, 11) if is_list else para(d, it, size=11))

    # ---- annexures (reproduced from the department template)
    for name in ("annex_wk.png", "annex_po.png"):
        page_break(d)
        picture(d, A / name, 6.2)

    # ---- abstract
    page_break(d)
    center(d, "ABSTRACT", 14, bold=True, before=6, after=14)
    para(d, "Game studios increasingly want to use **Generative AI (GenAI)** to speed up the design of characters, "
            "dialogue and story variations. Raw model output, however, is unreliable for production use: it can break "
            "the game's rules and lore, imitate copyrighted characters or artists, produce content unsuitable for the "
            "target age rating, consume significant compute budget, and weaken the creative control of human designers. "
            "This case study presents an **AI-assisted game-design workflow** that generates complete character "
            "packages—concept, abilities, backstory, concept-art portrait and image prompt, in-voice dialogue and "
            "branching story arcs—while enforcing these constraints. All generation is grounded in a machine-readable "
            "**Game Bible**, and every output passes through **input and output guardrails** (content filtering, "
            "prompt-injection blocking and copyright-similarity detection), a **rule-based consistency checker** "
            "(seven rules, scored out of 100) with **targeted auto-repair**, **cost controls** (token metering, caching, "
            "model routing and a budget cap) and a mandatory **human-in-the-loop review** with full version history. "
            f"In an evaluation of {len(ROWS) * ROWS[0]['runs']} generated characters, raw drafts at maximum creativity "
            f"violated at least one rule in {HI['drafts_with_violations_pct']}% of cases (average score "
            f"{HI['avg_draft_score']}), whereas the final outputs averaged **{HI['avg_final_score']:.0f}/100 with "
            f"{HI['final_with_errors_pct']}% residual errors**. The guardrails blocked {GUARD['unsafe_blocked']} unsafe "
            f"or injection briefs with {GUARD['safe_false_blocks']} false blocks and flagged "
            f"{GUARD['copycat_names_flagged']} look-alike names, and model routing reduced estimated cost per character "
            f"by about **{SAVING:.0f}×** compared with using the largest model for every task. The system runs fully "
            "offline as a Streamlit application and demonstrates how GenAI can accelerate creative work while keeping it "
            "consistent, safe, original, affordable and under human control.", size=11.5, line=1.6)

    # ---- acknowledgement
    page_break(d)
    center(d, "ACKNOWLEDGEMENT", 14, bold=True, before=6, after=14)
    ack = [
        "The satisfaction and euphoria that accompany the successful completion of any task would be incomplete without "
        "the mention of the people who made it possible, whose constant guidance and encouragement crowned the efforts "
        "with success.",
        "We would like to profoundly thank Management of **Dayananda Sagar College of Engineering** for providing such a "
        "healthy environment for the successful completion of the project work.",
        f"We would like to express our thanks to the Principal **{PRINCIPAL}** for their encouragement that motivated us "
        "for the successful completion of the project work.",
        f"We would wish to express our gratitude to the HoD, **{HOD}** and project coordinator **{GUIDE}, "
        f"{GUIDE_DESIG},** Department of Artificial Intelligence and Machine Learning for providing a good working "
        "environment and for her constant support and encouragement.",
        f"It gives me great pleasure in placing on record a deep sense of gratitude to our guide **{GUIDE}**, "
        f"{GUIDE_DESIG}, Department of Artificial Intelligence and Machine Learning for her expert guidance, initiative, "
        "encouragement and constant support throughout the course of project work.",
        "We would also like to thank all other teaching and non-teaching staff of the Artificial Intelligence and Machine "
        "Learning Department who have directly or indirectly helped me in the completion of this project work.",
        "We sincerely acknowledge and thank our parents for being a constant source of inspiration, encouragement, and "
        "support, which has been instrumental in the successful completion of this project work.",
    ]
    for a in ack:
        para(d, a, size=11.5, line=1.6, after=8)
    d.add_paragraph()
    para(d, f"**{STUDENT} ({USN})**", WD_ALIGN_PARAGRAPH.RIGHT, 11, line=1.2)


def contents(R, toc, figs, tabs):
    d = R.doc
    page_break(d)
    center(d, "TABLE OF CONTENTS", 14, bold=True, before=6, after=12)

    def row(label, page, level=0, bold=False, after=1):
        p = d.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, after, 1.25, indent=[0, 0.4, 0.8][level])
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.05), WD_TAB_ALIGNMENT.RIGHT)
        rich(p, f"{label}\t{page}", 11.5 if level else 12, bold=bold)

    row("DESCRIPTION", "PAGE NO.", bold=True, after=6)
    for lab, pg in [("Certificate", "i"), ("Declaration", "ii"), ("Abstract", "iii"), ("Acknowledgement", "iv"),
                    ("Contents", "v"), ("List of Figures", "vi"), ("List of Tables", "vii")]:
        row(lab, pg, bold=True)
    for level, label, key in toc:
        if level == 0:
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        row(label, R.pg(key) if key else "", level, bold=(level == 0))
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    row("References", R.pg("REFERENCES"), bold=True)

    for title, items, head in (("LIST OF FIGURES", figs, "FIGURE NO."), ("LIST OF TABLES", tabs, "TABLE NO.")):
        page_break(d)
        center(d, title, 14, bold=True, before=6, after=14)
        t = d.add_table(rows=1 + len(items), cols=3)
        borderless(t)
        widths = [1.4, 3.6, 1.2]
        fix_layout(t, widths)
        for i, r in enumerate([(head, "DESCRIPTION", "PAGE NO.")] + [(a, b, R.pg(k)) for a, b, k in items]):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.width = Inches(widths[j])
                p = c.paragraphs[0]
                fmt(p, WD_ALIGN_PARAGRAPH.CENTER if j == 2 else WD_ALIGN_PARAGRAPH.LEFT, 2, 4, 1.2)
                rich(p, v, 11.5 if i else 12, bold=(i == 0))


# ----------------------------------------------------------------------------- main body
def body(R):
    d = R.doc
    n_abil = sum(len(v) for v in ABILITY_POOL.values())
    n_lines = sum(len(lines) for p in PERSONALITIES.values() for lines in p.values())
    n_content = sum(len(v) for v in safety.CONTENT_RULES.values())

    # ======================================================== CHAPTER 1
    R.chapter(1, "Introduction", first=True)
    R.h2("1.1", "Introduction")
    para(d, "Character design is one of the most time-consuming parts of video-game development. Every playable hero, "
            "companion, villain and non-player character (NPC) needs a visual concept, abilities, a backstory, a voice "
            "and a place in the story, and all of it must remain consistent with the game's world and rules. "
            "**Generative AI (GenAI)**, in particular large language models (LLMs) and image models, can produce such "
            "content in seconds, which makes it attractive to studios that must populate large worlds with hundreds of "
            "characters.")
    para(d, "In this case study, an **AI-assisted Character Design Studio** is developed for a fictional high-fantasy "
            f"role-playing game, **{BIBLE['world_name']}**. The system takes a short designer brief (role, faction, class, "
            "personality and optional notes) and generates a complete character package: name, appearance, colour "
            "palette, abilities, signature item, backstory, motivation, flaw, catchphrase, a concept-art portrait, a prompt "
            "for an external image model, in-voice dialogue for different game situations and several branching story "
            "arcs.")
    para(d, "Raw GenAI output, however, cannot be trusted directly. It may use the wrong magic for a faction, introduce "
            "technology that does not exist in a medieval world, contradict established lore, reuse names of famous "
            "characters, imitate a living artist's style, or include content that violates the game's age rating. "
            "Generation also has a real **computational cost**, and studios must keep **human designers in creative "
            "control**. The proposed system therefore wraps the generator in a **guarded pipeline** consisting of a "
            "machine-readable **Game Bible**, **input and output guardrails**, a **consistency checker with automatic "
            "repair**, **cost controls** and a mandatory **human review** step. The overall objective is to show that "
            "GenAI can accelerate creative work while remaining consistent, safe, original, affordable and "
            "human-led.")
    R.h3("1.1.1", "Background")
    para(d, "Game studios have long used **procedural content generation (PCG)**—rule-based or search-based algorithms—to "
            "create levels, items and characters automatically. PCG guarantees that output respects the designer's rules, "
            "but its variety is limited by hand-written templates. Modern LLMs, in contrast, produce fluent and varied "
            "text from a natural-language prompt, but they offer no guarantee of consistency, safety or originality.")
    para(d, "This case study combines the strengths of both: the generator is free to be creative, while explicit rules "
            "derived from the Game Bible are used to **verify, score and repair** every output. The same pipeline "
            "works with a built-in **offline procedural engine** (no internet or API key required) and with a free "
            "**local open-source LLM** through Ollama, so the guardrails are independent of the particular model used.")
    R.h3("1.1.2", "Motivation")
    para(d, "A single open-world game can contain hundreds of named characters, each requiring several days of work from "
            "writers and concept artists. GenAI promises to reduce this effort dramatically, but uncontrolled use creates "
            "new problems: **lore inconsistencies** that break immersion, **intellectual-property (IP) risk** from "
            "look-alike characters, **age-rating violations** that can delay certification, **unpredictable cloud "
            "costs**, and concern among creative staff that AI will replace rather than support them.")
    para(d, "This motivates the development of a workflow in which GenAI is used as a **creative assistant** rather than "
            "an unchecked author: the AI drafts, automated checks catch and repair its mistakes, and a human designer "
            "makes every final decision. Such a workflow lets studios benefit from AI speed while addressing the "
            "constraints explicitly required by the case study: **copyright, inappropriate content, computational cost "
            "and human creative control.**")

    R.h2("1.2", "Literature Survey")
    para(d, "Research relevant to this case study spans procedural content generation in games, large language models "
            "for narrative generation, grounding and self-correction of model output, safety and prompt-injection "
            "attacks, memorisation and copyright, and cost-efficient use of LLMs.")
    R.h3("1.2.1", "Literature review")
    para(d, "Togelius et al. [1] surveyed **search-based procedural content generation**, in which candidate content is "
            "generated and then evaluated against fitness functions; this generate-and-test idea directly inspires the "
            "generate–check–repair loop used in this work. Summerville et al. [2] reviewed **procedural content generation "
            "via machine learning (PCGML)**, noting that learned generators can produce more varied content but provide "
            "weaker guarantees and may reproduce their training data. Yannakakis and Togelius [3] describe the broader role "
            "of AI in games, including content generation and the importance of designer control.")
    para(d, "Brown et al. [4] showed that large language models can perform new generation tasks from only a prompt and a "
            "few examples, which makes them suitable for dialogue and backstory writing without task-specific training. "
            "Park et al. [5] used LLMs to drive **believable generative agents** with memory and planning, demonstrating "
            "both the potential of AI-driven characters and the difficulty of keeping their behaviour consistent over "
            "time.")
    para(d, "To improve factual consistency, Lewis et al. [6] proposed **retrieval-augmented generation (RAG)**, which "
            "grounds generation in an external knowledge source; in this project the Game Bible plays that role. Madaan "
            "et al. [7] introduced **Self-Refine**, in which a model iteratively improves its output using feedback; the "
            "proposed auto-repair step applies the same idea but uses deterministic rule feedback and regenerates only "
            "the fields that failed. Bai et al. [8] used an explicit set of principles to steer models towards harmless "
            "behaviour, and Ouyang et al. [12] showed the value of **human feedback** in aligning model outputs with "
            "user intent—both motivating explicit rules and a human reviewer in this system.")
    para(d, "On the risk side, Perez and Ribeiro [9] demonstrated **prompt-injection** attacks such as “ignore previous "
            "instructions”, which the input guardrail here detects. Carlini et al. [10] showed that generative image "
            "models can **memorise and reproduce training images**, highlighting the copyright risk of imitating existing "
            "characters and styles. The OWASP Top 10 for LLM Applications [13] lists prompt injection, sensitive "
            "output and unbounded consumption among the main risks of LLM systems. Finally, Chen et al. [11] proposed "
            "**FrugalGPT**, showing that routing requests between cheaper and more expensive models can substantially "
            "reduce cost without sacrificing quality, which informs the model-routing and caching strategy used here.")
    para(d, "Based on these studies, the present work combines **Game-Bible grounding, rule-based verification with "
            "targeted self-repair, input/output guardrails, cost-aware routing and human-in-the-loop review** into a "
            "single, practical workflow for game character design.")
    R.h3("1.2.2", "Problem Statement")
    para(d, "A game-development studio wants to use GenAI to generate character concepts, dialogue and story variations. "
            "Unconstrained generation frequently violates game rules and narrative constraints, may infringe copyright, "
            "may produce inappropriate content, incurs computational cost, and can reduce the creative control of "
            "designers.")
    para(d, "Therefore, the problem is to develop an **AI-assisted game-design workflow** that generates character "
            "concepts, dialogue and story variations **while maintaining consistency with game rules and narrative "
            "constraints**, and that explicitly addresses **copyright, inappropriate content, computational cost and "
            "human creative control.**")
    R.h3("1.2.3", "Stakeholders")
    para(d, "The main stakeholders involved in the Character Design Studio are:", after=2)
    for b in ["**Game Studios and Producers:** Want faster content creation at a predictable cost and quality.",
              "**Narrative Designers and Concept Artists:** Use the AI drafts and keep final creative authority.",
              "**Players:** Expect a consistent world and content appropriate to the game's age rating.",
              "**IP Holders and Legal Teams:** Need assurance that characters do not imitate protected works.",
              "**Rating Bodies (e.g., PEGI [14]):** Assess whether the game's content suits its target age group.",
              "**Developers and ML Engineers:** Build, monitor and improve the generation pipeline."]:
        bullet(d, b)
    R.h3("1.2.4", "Constraints and Assumptions")
    para(d, "**Constraints:**", after=2, keep=True)
    for b in ["The system must run on a student laptop without paid cloud APIs or API keys.",
              "Content must respect the Game Bible: faction magic, tech level, lore, palettes and naming conventions.",
              f"Content must be suitable for the **{BIBLE['age_rating']}** age rating.",
              "Copyright checks use a curated sample list of protected names, franchise terms and artist styles.",
              "Generation must stay within a configurable per-session budget."]:
        bullet(d, b)
    para(d, "**Assumptions:**", after=2, before=4, keep=True)
    for b in ["The Game Bible is complete and correct; it is the reference against which consistency is judged.",
              "A human designer reviews every character before it is used in the game.",
              "Token usage is estimated at roughly four characters per token, and model prices are illustrative tiers.",
              "The offline engine's deliberate “drift” reproduces the typical mistakes of real LLMs, so the guardrails can "
              "be tested reliably and repeatably."]:
        bullet(d, b)

    R.h2("1.3", "Proposed System")
    para(d, "The proposed system is an **AI-assisted Character Design Studio** implemented as a web application. A designer "
            "submits a brief, and the request flows through a **guarded pipeline**: (i) an input guardrail screens the "
            "brief for unsafe content, prompt injection and copyright issues; (ii) a cache, budget check and model router "
            "control cost; (iii) the generator produces a character grounded in the Game Bible; (iv) an output guardrail "
            "checks safety and copyright; (v) a consistency checker applies seven rules and computes a score out of 100; "
            "(vi) if errors are found, an **auto-repair loop** regenerates only the broken fields; and (vii) the result is "
            "saved for **human designer review**.")
    para(d, "Approved characters can be extended with dialogue for six game situations and five types of story arc, "
            "given one of several generated concept-art portraits (or the designer's own artwork), and exported as JSON "
            "for a game engine or as a PDF character sheet that records the full provenance of every AI and human change.")

    R.h2("1.4", "Data Collection")
    para(d, "Unlike a classical machine-learning project, this system is driven mainly by **structured domain knowledge** "
            "rather than a large training dataset. The knowledge sources were authored and curated for this case study "
            "and are summarised in Table 1.1. Separate evaluation sets were prepared to measure the guardrails.")
    R.tab("1.1", "Knowledge Sources and Evaluation Data")
    table(d, ["Source", "Contents", "Used for"], [
        ["Game Bible (JSON)", f"World, tone, {BIBLE['age_rating']} rating, art style, tech level, "
                              f"{len(BIBLE['forbidden_tech'])} forbidden-technology terms, {len(BIBLE['lore_facts'])} lore "
                              f"facts, {len(BIBLE['lore_contradictions'])} lore-contradiction patterns, "
                              f"{len(BIBLE['factions'])} factions, {len(BIBLE['classes'])} classes, "
                              f"{len(BIBLE['roles'])} roles", "Prompt grounding and consistency rules"],
        ["Generation vocabulary", f"{n_abil} elemental ability templates, {len(PERSONALITIES)} personality voices with "
                                  f"{n_lines} dialogue lines across {len(SITUATIONS)} situations, {len(ARC_TEMPLATES)} "
                                  "story-arc templates", "Offline generation of characters, dialogue and arcs"],
        ["Content-safety lexicon", f"{len(safety.CONTENT_RULES)} categories, {n_content} terms; "
                                   f"{len(safety.INJECTION_PATTERNS)} prompt-injection patterns", "Input and output guardrails"],
        ["Protected-IP lists", f"{len(safety.PROTECTED_NAMES)} protected character names, {len(safety.FRANCHISE_TERMS)} "
                               f"franchise terms, {len(safety.SIGNATURE_TRAITS)} signature-trait combinations, artist/"
                               "studio style patterns", "Copyright-similarity detection"],
        ["Evaluation sets", f"{len(ROWS) * ROWS[0]['runs']} generated characters (4 creativity levels), 10 unsafe/injection "
                            "briefs, 10 safe briefs, 10 copycat names, 132 original generated names",
         "Measuring consistency, repair and guardrail accuracy"],
    ], [1.55, 3.35, 1.4], size=10.5)
    R.h3("1.4.1", "Data Cleaning and Preprocessing")
    para(d, "All designer input and model output is **normalised** before checking: text is lower-cased and tokenised "
            "with regular expressions so that rules match whole words (for example, “scar” must not match “car”). For "
            "name comparison, non-letter characters are removed before computing similarity. The Game Bible is validated "
            "as JSON before it is saved, and output from the local LLM is parsed as JSON and checked against the "
            "expected character schema; any missing or malformed field is filled by the offline engine. User-supplied "
            "text is HTML-escaped before display, and blocked terms can be **redacted** so that flagged output is never "
            "shown in raw form.")
    R.h3("1.4.2", "Feature Selection and Scaling")
    para(d, "Each character is described by a fixed set of **features**: role, faction, class, personality and pronouns "
            "(inputs), and name, palette, abilities with their elements, signature item, appearance, backstory, "
            "motivation, flaw, catchphrase and image prompt (outputs). Ability text is mapped to magic **elements** "
            "using keyword groups (ice, fire, nature, earth, water, storm), which allows the checker to compare each "
            "ability with the faction's allowed elements.")
    para(d, "Rule results are converted to a single **consistency score on a 0–100 scale**: each error deducts 20 "
            "points and each warning deducts 7. Copyright similarity is measured with a normalised sequence-matching "
            "ratio between 0 and 1, and names scoring **0.8 or higher** are flagged. The designer's **creativity "
            "(temperature)** setting is scaled to a drift probability between 0% and 60% per field, and token usage is "
            "estimated from text length to calculate cost.")
    R.h3("1.4.3", "Tools and Techniques Used")
    para(d, "The following tools and techniques were used for developing the Character Design Studio:", after=2)
    for b in ["**Python 3.12:** Core implementation of the pipeline, rules and evaluation.",
              "**Streamlit:** Interactive web interface with seven workspaces and a custom theme.",
              "**Ollama (optional):** Runs a free, local open-source LLM (e.g., Llama 3.2) without API keys.",
              "**Pandas & Matplotlib:** Cost analysis, evaluation tables and charts.",
              "**difflib and regular expressions:** Name-similarity scoring and rule matching.",
              "**fpdf2 & python-pptx:** PDF character sheets, report and presentation generation.",
              "**pytest & Streamlit AppTest:** Nine automated tests, including a headless test of the full UI.",
              "**Techniques:** Knowledge grounding, rule-based verification, targeted self-repair, content moderation, "
              "prompt-injection detection, similarity search, response caching, model routing and human-in-the-loop "
              "review."]:
        bullet(d, b)

    # ======================================================== CHAPTER 2
    R.chapter(2, "Evaluation of Alternatives and Decision Making")
    R.h2("2.1", "Architectures Considered")
    para(d, "Four approaches were considered for the AI-assisted character design workflow:", after=2)
    for i, b in enumerate([
            "**Direct LLM prompting:** The designer's brief is sent straight to a language model and the output is used as "
            "it is. This is simple and creative, but provides no guarantee of consistency, safety or originality.",
            "**Fine-tuned model:** A model is fine-tuned on the studio's existing lore and characters. This improves style "
            "and consistency, but needs a large dataset and GPU training, must be retrained whenever the Game Bible "
            "changes, and may memorise and reproduce existing content.",
            "**Template-based procedural generation:** Characters are assembled only from hand-written templates and "
            "rules. Output is always valid and free, but variety and creativity are limited.",
            "**Grounded generation with guardrails, auto-repair and human review (proposed):** A creative generator is "
            "grounded in the Game Bible, and every output is screened, scored, repaired where necessary and approved by a "
            "human designer."], 1):
        numbered(d, i, b)
    para(d, "All approaches were compared on the same design constraints: rule consistency, creativity, safety and "
            "copyright control, cost, human control and practical feasibility. Fig 2.1 shows the architecture of the "
            "proposed approach.", before=4)
    R.fig(A / "fig_architecture.png", 4.9, "2.1", "Proposed Architecture")

    R.h2("2.2", "Performance Metrics")
    para(d, "The performance of the workflow was evaluated using the following metrics:", after=2)
    for b in ["**Consistency score:** 100 minus 20 points per rule error and 7 points per warning, averaged over characters.",
              "**Draft violation rate:** Percentage of raw drafts that break at least one rule or guardrail.",
              "**Residual error rate:** Percentage of final characters that still contain a rule error after repair.",
              "**Guardrail block rate:** Percentage of unsafe or prompt-injection briefs that are blocked.",
              "**False block rate:** Percentage of safe briefs that are wrongly blocked.",
              "**Copyright flag rate and false flags:** Look-alike names detected, and original names wrongly flagged.",
              "**Cost per character:** Estimated token cost with routing and caching, compared with using the largest "
              "model for every task.",
              "**Repair rounds:** Average number of automatic repair iterations needed per character."]:
        bullet(d, b)
    para(d, "Using several metrics, instead of output quality alone, gives a balanced view of consistency, safety, cost "
            "and the amount of correction the AI requires.", before=4)

    R.h2("2.3", "Comparative Analysis")
    para(d, "Table 2.1 compares the four approaches qualitatively against the case-study constraints.")
    R.tab("2.1", "Comparison of Architectures")
    table(d, ["Criterion", "Direct LLM prompting", "Fine-tuned model", "Template PCG", "Proposed pipeline"], [
        ["Rule consistency", "Low", "Medium", "High", "High (checked + repaired)"],
        ["Creativity / variety", "High", "High", "Low", "High"],
        ["Safety & copyright", "None", "Low (may memorise)", "High", "High (input + output guardrails)"],
        ["Computational cost", "Medium–High", "High (training)", "Very low", "Low (routing + caching)"],
        ["Human control", "Low", "Low", "Medium", "High (mandatory review)"],
        ["Adapts to rule changes", "Partial", "Needs retraining", "Manual edits", "Immediate (edit Game Bible)"],
    ], [1.45, 1.2, 1.2, 1.05, 1.5], size=10, align_cols="lccc".ljust(5, "c"), bold_first=True)
    para(d, "To measure the benefit of the guardrails quantitatively, the **raw drafts** (equivalent to direct, unguarded "
            f"generation) were compared with the **final outputs** of the proposed pipeline over {ROWS[0]['runs']} random "
            "briefs at each creativity level. The results are shown in Table 2.2 and Fig 2.2.")
    R.tab("2.2", "Unguarded Drafts vs. Proposed Pipeline")
    table(d, ["Creativity", "Drafts with violations", "Avg. draft score", "Avg. final score", "Final with errors",
              "Repair rounds"],
          [[r["temperature"], f"{r['drafts_with_violations_pct']}%", r["avg_draft_score"], r["avg_final_score"],
            f"{r['final_with_errors_pct']}%", r["avg_repair_rounds"]] for r in ROWS],
          [1.0, 1.15, 1.0, 1.0, 1.05, 1.0], size=10.5, align_cols="cccccc")
    R.fig(A / "fig_scores.png", 5.6, "2.2", "Consistency Score Before and After Auto-Repair")
    para(d, f"As creativity increases, raw drafts break the rules more often (up to {HI['drafts_with_violations_pct']}% "
            f"of drafts, with an average score of {HI['avg_draft_score']}), while the proposed pipeline keeps the final "
            f"average at {HI['avg_final_score']:.0f} with {HI['final_with_errors_pct']}% residual errors. The guardrails "
            "were evaluated separately on dedicated test sets (Table 2.3).")
    R.tab("2.3", "Guardrail Accuracy")
    table(d, ["Test", "Result"], [
        ["Unsafe or prompt-injection briefs blocked", GUARD["unsafe_blocked"]],
        ["Safe briefs wrongly blocked", GUARD["safe_false_blocks"]],
        ["Copycat names flagged (e.g., Pikachoo, Kratoz)", GUARD["copycat_names_flagged"]],
        ["Original generated names wrongly flagged", GUARD["original_names_false_flags"]],
    ], [4.2, 1.6], size=10.5, align_cols="lc")
    para(d, "The single borderline flag (“Maro”, 89% similar to “Mario”) shows why copyright results are presented as "
            "**flags for a human to judge** rather than automatic blocks. Fig 2.3 compares the estimated cost per "
            "character of the routed pipeline with using the largest model for every task.")
    R.fig(A / "fig_cost.png", 5.6, "2.3", "Estimated Cost per Character")

    R.h2("2.4", "Justification of Selected Solution")
    para(d, "Based on the comparative analysis, the **grounded generation pipeline with guardrails, auto-repair and human "
            "review** was selected. It is the only approach that keeps the **creativity of GenAI** while meeting every "
            f"constraint of the case study: final outputs averaged **{HI['avg_final_score']:.0f}/100** consistency even "
            f"when {HI['drafts_with_violations_pct']}% of raw drafts broke the rules, all unsafe and injection briefs in "
            "the test set were blocked without false blocks, every copycat name was flagged, and estimated cost was about "
            f"**{SAVING:.0f}× lower** than using the largest model for everything.")
    para(d, "Direct prompting was rejected because it provides no safety or consistency guarantees; fine-tuning was "
            "rejected because it requires large datasets and retraining and may memorise protected content; and "
            "template-only generation was rejected because it cannot provide sufficient creative variety. Because the "
            "guardrails are independent of the generator, the selected design can later use a stronger model without "
            "changing the rules or the review process.")

    # ======================================================== CHAPTER 3
    R.chapter(3, "Solution and Impacts")
    R.h2("3.1", "Proposed Solution")
    para(d, "The proposed solution is a **GenAI Character Design Studio** delivered as a Streamlit web application with "
            "seven workspaces: Overview, Game Bible, Generate, Dialogue & Story, Designer Review, Cost and Export. "
            "Every request follows the guarded pipeline of Fig 2.1, and each stage is shown to the designer as a "
            "step-by-step trace (Fig 3.1).")
    R.fig(S / "03_generate_result.png", 5.9, "3.1", "Character Generation Interface")
    para(d, "When the AI makes mistakes, the system explains exactly what was wrong and how it was repaired. In the example "
            "of Fig 3.2, a single raw draft contained forbidden technology (a laser), a lore contradiction (a tamed "
            "dragon, although dragons are extinct), a name copied from a famous character and an imitation of a real "
            "studio's art style; the checker caught every issue and auto-repair raised the score from 53 to 100.")
    R.fig(S / "04_self_correction.png", 5.0, "3.2", "Automatic Detection and Repair of Rule Violations")
    para(d, "The system is intended to **support designers**, not replace them: nothing can be exported until a human "
            "designer approves it, and every AI and human change is versioned.")

    R.h2("3.2", "Implementation Strategy")
    para(d, "The Character Design Studio was implemented in **Python** using **Streamlit**. The implementation was "
            "carried out in the following stages:", after=2)
    for i, b in enumerate([
            "**Game Bible design:** Defined the world, factions, magic elements, lore, tech level, palettes, naming "
            "conventions and banned content as structured JSON.",
            "**Generator:** Implemented an offline procedural engine with a creativity-controlled drift that imitates LLM "
            "errors, and an optional Ollama engine for a local open-source LLM, behind a common interface.",
            "**Guardrails:** Implemented age-rating content filters, prompt-injection detection and copyright-similarity "
            "checks for both input and output (Fig 3.3).",
            "**Consistency checker and auto-repair:** Implemented rules R1–R7 with error/warning severities, a 0–100 "
            "score and a repair loop that regenerates only the failed fields.",
            "**Cost control:** Added token metering, response caching, task-based model routing and a session budget cap.",
            "**Human review and provenance:** Added approve/reject/edit/regenerate actions, version history with diffs, "
            "and concept-art selection or upload (Fig 3.4).",
            "**Evaluation and testing:** Measured consistency, guardrail accuracy and cost, and wrote nine automated "
            "tests, including a headless test of the complete user interface."], 1):
        numbered(d, i, b)
    R.fig(S / "08_guardrail_block.png", 5.9, "3.3", "Input Guardrail Blocking an Unsafe Brief")
    R.fig(S / "07_concept_art.png", 5.9, "3.4", "Designer Review with Concept-Art Selection")

    R.h2("3.3", "Sector-wide Impact")
    para(d, "The proposed workflow can benefit the **video-game and interactive-media sector** by reducing the time needed "
            "to create first drafts of characters, dialogue and quests, especially for small and independent studios "
            "with limited writing and art staff. Because outputs are grounded in a shared Game Bible and checked "
            "automatically, large teams can generate content in parallel without breaking the consistency of the world.")
    para(d, "The same pattern—**grounded generation, automated verification, guardrails, cost control and human "
            "approval**—can be applied to other creative and knowledge domains such as animation, educational content, "
            "marketing and technical documentation. However, practical adoption requires continued attention to "
            "**intellectual-property rights, content safety, fair treatment of creative workers and transparency** "
            "about AI involvement. The system should therefore be used as a decision-support and drafting tool rather "
            "than a replacement for human creativity.")

    R.h2("3.4", "Sustainability Development Goals Mapping")
    para(d, "The proposed Character Design Studio contributes to the following **United Nations Sustainable Development "
            "Goals (SDGs):**", after=2)
    for b in ["**SDG 8 – Decent Work and Economic Growth:** Augments rather than replaces creative professionals, keeping "
              "designers in control while improving productivity.",
              "**SDG 9 – Industry, Innovation and Infrastructure:** Promotes responsible adoption of GenAI in the creative "
              "and software industries.",
              "**SDG 12 – Responsible Consumption and Production:** Reduces unnecessary computation through caching, model "
              "routing and targeted repair, lowering the energy cost of AI use.",
              "**SDG 16 – Peace, Justice and Strong Institutions:** Supports respect for intellectual-property rights and "
              "accountable decision-making through provenance records."]:
        bullet(d, b)
    para(d, "Overall, the system demonstrates how **AI-assisted creative tools can support innovation while remaining "
            "responsible, efficient and human-centred.**", before=4)

    R.h2("3.5", "Limitations and Future Improvements")
    para(d, "**Limitations**", after=2, keep=True)
    for b in ["Keyword- and pattern-based filters can miss paraphrased unsafe content or indirect references.",
              "Copyright checks use a sample list of protected names and styles rather than a complete IP database.",
              "The default offline engine simulates LLM behaviour; results with a full-scale LLM may differ.",
              "Concept-art portraits are illustrative vector drawings rather than production artwork.",
              "Cost figures are estimates based on illustrative price tiers and approximate token counts."]:
        bullet(d, b)
    para(d, "**Future Improvements**", before=6, after=2, keep=True)
    for b in ["Add a trained **moderation model** and **embedding-based similarity search** for stronger safety and IP checks.",
              "Connect to a **licensed IP database** for comprehensive copyright screening.",
              "Integrate a **local image-generation model** with the same style guardrails for concept art.",
              "Support **multi-role review** (writer, art director, legal) with comments and approval workflows.",
              "Fine-tune a small model on approved characters to capture the studio's house style.",
              "Evaluate the workflow with professional designers and real game projects."]:
        bullet(d, b)

    R.h2("3.6", "Mapping of Course Outcome and Program Outcome")
    R.tab("3.1", "Course Outcome Mapping")
    table(d, ["Course Outcome", "Mapping", "Justification"], [
        ["CO1", "Yes", "Applies GenAI and AI-safety techniques to the real-world problem of game character design."],
        ["CO2", "Yes", "Involves collection and curation of the Game Bible and IP/safety lexicons, text normalisation, "
                       "feature definition and score scaling."],
        ["CO3", "Yes", "Compares four alternative architectures using consistency, safety, copyright and cost metrics."],
        ["CO4", "Partially", "The project does not use IoT, but it addresses intelligent automation, responsible AI and "
                             "sector-level impact."],
    ], [1.3, 1.0, 4.0], size=11, align_cols="ccl")
    R.tab("3.2", "Program Outcome Mapping")
    table(d, ["Program Outcome", "Mapping", "Justification"], [
        ["PO1 - Engineering Knowledge", "Yes", "Applies computing, AI and software-engineering concepts to controlled "
                                               "content generation."],
        ["PO2 - Problem Analysis", "Yes", "Analyses rule violations, safety and copyright risks, and cost drivers of "
                                          "GenAI output."],
        ["PO3 - Design/Development of Solutions", "Yes", "Designs a guarded pipeline that meets consistency, safety, cost "
                                                         "and human-control requirements."],
        ["PO4 - Conduct Investigations", "Yes", "Evaluates draft vs. final quality, guardrail accuracy and cost across "
                                                "creativity levels."],
        ["PO5 - Engineering Tool Usage", "Yes", "Uses Python, Streamlit, Ollama, Pandas, Matplotlib and pytest."],
        ["PO7 - Ethics", "Yes", "Addresses intellectual property, age-appropriate content, transparency and the "
                                "role of human creators."],
    ], [1.9, 0.9, 3.5], size=11, align_cols="ccl")

    R.h2("3.7", "Blooms and TRL level of the work")
    para(d, "**Bloom’s Taxonomy**", after=2, keep=True)
    para(d, "The project involves the following levels of Bloom’s Taxonomy:", after=2)
    R.tab("3.3", "Blooms Taxonomy")
    blooms(d)
    para(d, "**Technology Readiness Level (TRL)**", before=8, after=2, keep=True)
    R.tab("3.4", "TRL Mapping")
    table(d, ["Parameter", "Level", "Justification"], [
        ["Technology Readiness Level (TRL)", "TRL 4", "The proposed system has been implemented as a working application "
                                                      "and validated with automated tests and a quantitative evaluation in "
                                                      "a controlled laboratory environment. It has not yet been deployed "
                                                      "in a commercial game studio."],
    ], [1.7, 0.9, 3.7], size=11, align_cols="ccl", bold_first=True)
    para(d, "The Character Design Studio has been implemented and tested on a laptop with an offline engine and an "
            "optional local LLM. Further evaluation with professional designers, full-scale models, real game projects "
            "and production infrastructure would be required to reach higher TRL levels.")

    R.h2("3.8", "Rubrics")
    R.tab("3.5", "Rubrics")
    rubrics(d)

    R.h2("3.9", "Plagiarism Report")
    para(d, "[Attach the plagiarism report here.]", WD_ALIGN_PARAGRAPH.CENTER, 11, italic=True, before=40)

    # ======================================================== REFERENCES
    page_break(d)
    center(d, "REFERENCES", 14, bold=True, before=6, after=12)
    R.toc.append((-1, "References", "REFERENCES"))
    refs = [
        "J. Togelius, G. N. Yannakakis, K. O. Stanley, and C. Browne, “Search-Based Procedural Content Generation: A "
        "Taxonomy and Survey,” IEEE Transactions on Computational Intelligence and AI in Games, vol. 3, no. 3, pp. "
        "172–186, 2011.",
        "A. Summerville, S. Snodgrass, M. Guzdial, C. Holmgård, A. K. Hoover, A. Isaksen, A. Nealen, and J. Togelius, "
        "“Procedural Content Generation via Machine Learning (PCGML),” IEEE Transactions on Games, vol. 10, no. 3, pp. "
        "257–270, 2018.",
        "G. N. Yannakakis and J. Togelius, Artificial Intelligence and Games. Cham, Switzerland: Springer, 2018.",
        "T. B. Brown et al., “Language Models are Few-Shot Learners,” in Advances in Neural Information Processing "
        "Systems (NeurIPS), vol. 33, 2020.",
        "J. S. Park, J. C. O’Brien, C. J. Cai, M. R. Morris, P. Liang, and M. S. Bernstein, “Generative Agents: "
        "Interactive Simulacra of Human Behavior,” in Proc. ACM Symposium on User Interface Software and Technology "
        "(UIST), 2023.",
        "P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” in Advances in Neural "
        "Information Processing Systems (NeurIPS), vol. 33, 2020.",
        "A. Madaan et al., “Self-Refine: Iterative Refinement with Self-Feedback,” in Advances in Neural Information "
        "Processing Systems (NeurIPS), vol. 36, 2023.",
        "Y. Bai et al., “Constitutional AI: Harmlessness from AI Feedback,” arXiv:2212.08073, 2022.",
        "F. Perez and I. Ribeiro, “Ignore Previous Prompt: Attack Techniques for Language Models,” NeurIPS ML Safety "
        "Workshop, arXiv:2211.09527, 2022.",
        "N. Carlini, J. Hayes, M. Nasr, M. Jagielski, V. Sehwag, F. Tramèr, B. Balle, D. Ippolito, and E. Wallace, "
        "“Extracting Training Data from Diffusion Models,” in Proc. 32nd USENIX Security Symposium, 2023.",
        "L. Chen, M. Zaharia, and J. Zou, “FrugalGPT: How to Use Large Language Models While Reducing Cost and Improving "
        "Performance,” arXiv:2305.05176, 2023.",
        "L. Ouyang et al., “Training Language Models to Follow Instructions with Human Feedback,” in Advances in Neural "
        "Information Processing Systems (NeurIPS), vol. 35, 2022.",
        "OWASP Foundation, “OWASP Top 10 for Large Language Model Applications,” 2025. [Online]. Available: "
        "https://genai.owasp.org",
        "Pan European Game Information (PEGI), “PEGI Age Ratings.” [Online]. Available: https://pegi.info",
    ]
    for i, r in enumerate(refs, 1):
        p = d.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 1.25, indent=0.4, first=-0.4)
        rich(p, f"[{i}]  {r}", 11)


def blooms(d):
    hdr_wk = [str(i) for i in range(1, 10)]
    hdr_po = [str(i) for i in range(1, 12)]
    hdr_pso = ["1", "2", "3"]
    ncol = 1 + 9 + 11 + 3
    rows = {
        "1": ({2, 3}, {"1": "-"}, {1, 5, 8, 9, 11}),
        "2": ({2, 3, 4, 5}, {"1": "-"}, {2, 4, 5, 8, 9, 11}),
        "3": ({2, 3}, {"1": "-"}, {3, 4, 5, 8, 9, 11}),
        "4": ({1, 2, 3, 4, 5, 8, 9}, {}, {3, 5, 8, 9, 11}),
    }
    overall_po = ["2", "3", "2", "3", "2", "2", "1", "3", "1", "2", "1"]
    overall_pso = ["3", "2", ""]
    t = d.add_table(rows=2 + len(rows) + 1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [0.95] + [0.205] * 9 + [0.225] * 11 + [0.215] * 3

    def put(i, j, txt, bold=False, size=9):
        c = t.cell(i, j)
        p = c.paragraphs[0]
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 1, 1, 1.0)
        rich(p, txt, size, bold=bold)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    merge(t.cell(0, 0), t.cell(1, 0))
    put(0, 0, "Course Outcomes", True)
    merge(t.cell(0, 1), t.cell(0, 9))
    put(0, 1, "Washington Accord Knowledge and Attitude Profile (WKs)", True)
    merge(t.cell(0, 10), t.cell(0, 20))
    put(0, 10, "Program Outcomes", True)
    merge(t.cell(0, 21), t.cell(0, 23))
    put(0, 21, "PSO", True)
    for j, h in enumerate(hdr_wk + hdr_po + hdr_pso, 1):
        put(1, j, h, True)
    for i, (co, (wks, special, pos)) in enumerate(rows.items(), 2):
        put(i, 0, co, True, 10)
        for k in range(1, 10):
            put(i, k, special.get(str(k), "✓" if k in wks else ""))
        for k in range(1, 12):
            put(i, 9 + k, "3" if k in pos else "")
        for k in range(3):
            put(i, 21 + k, "3")
    last = 2 + len(rows)
    put(last, 0, "Overall", True, 10)
    for k, v in enumerate(overall_po):
        put(last, 10 + k, v)
    for k, v in enumerate(overall_pso):
        put(last, 21 + k, v)
    fix_layout(t, widths)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def rubrics(d):
    rows = [
        ["Problem Analysis", "CO1", "PO1, PO5, PO8, PO9, PO11",
         "Clearly identifies the complex engineering problem, stakeholders, constraints, assumptions and root causes using "
         "systematic analysis.", "Identifies the problem and most influencing factors with minor gaps.",
         "Basic understanding of the problem with limited analysis.", "Problem analysis is incomplete or incorrect.", ""],
        ["Application of Concept Knowledge", "CO1, CO2", "PO1, PO2, PO4, PO5, PO8, PO9, PO11",
         "Effectively applies relevant AI/ML concepts, algorithms, tools and data-preparation techniques (collection, "
         "cleaning, feature selection, scaling) to analyse the case.",
         "Applies appropriate concepts and data-preparation techniques with minor errors.",
         "Applies only basic concepts with limited relevance.", "Inappropriate or no application of AI/ML concepts.", ""],
        ["Evaluation of Alternatives & Decision Making", "CO3", "PO3, PO4, PO5, PO8, PO9, PO11",
         "Critically evaluates and compares multiple ML models/deep-learning architectures using appropriate performance "
         "metrics, considering feasibility, cost, scalability, ethics and risk before selecting the best solution.",
         "Evaluates more than one model/solution with reasonable justification.",
         "Considers limited alternatives with basic justification.", "No meaningful comparison or justification provided.", ""],
        ["Solution & Impact Assessment", "CO4", "PO3, PO5, PO8, PO9, PO11",
         "Critically assesses the proposed AI (incl. AI+IoT, where relevant) solution's integration and impact across the "
         "target sector, explicitly linking expected outcomes to relevant sustainability and societal-development goals "
         "(e.g., environmental, economic, health, or social well-being impact), with clear implementation strategy, "
         "limitations and future improvements.",
         "Proposes a feasible solution with adequate impact explanation and reasonable, if partial, consideration of its "
         "sustainability or societal-development impact.",
         "Provides a basic solution with limited impact analysis and only a passing or generic mention of "
         "sustainability/societal impact.",
         "Solution is incomplete, impractical, or unsupported, and shows no consideration of sustainability or "
         "societal-development impact.", ""],
        ["Presentation & Technical Communication", "CO1–CO4", "PO9",
         "Presents the case logically with supporting evidence (data, diagrams, references), communicates confidently, and "
         "responds effectively to questions.", "Well-organized presentation with minor omissions.",
         "Average presentation with limited supporting evidence.", "Poor organization and unable to justify the analysis.", ""],
        ["Total", "", "", "", "", "", "", "20"],
    ]
    t = table(d, ["Criteria", "CO", "PO", "Excellent (4 Marks)", "Good (3 Marks)", "Satisfactory (2 Marks)",
                  "Needs Improvement (0–1 Mark)", "Marks"], rows,
              [0.8, 0.5, 0.65, 1.55, 0.95, 0.95, 0.85, 0.5], size=8, align_cols="lccllllc", bold_first=True)
    last = t.rows[-1].cells
    m = merge(last[0], last[6])
    m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ----------------------------------------------------------------------------- header / footer
def main_section(R):
    d = R.doc
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    # remove page border inherited from section 1
    pb = sec._sectPr.find(qn("w:pgBorders"))
    if pb is not None:
        sec._sectPr.remove(pb)
    restart_numbering(sec, 1)
    hp = sec.header.paragraphs[0]
    fmt(hp, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0)
    rich(hp, SHORT_TITLE, 9, name="Calibri")
    para_border(hp, "bottom")
    fp = sec.footer.paragraphs[0]
    fmt(fp, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0)
    para_border(fp, "top", val="thickThinSmallGap")
    ts = fp.paragraph_format.tab_stops
    ts.add_tab_stop(Inches(3.17), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Inches(6.35), WD_TAB_ALIGNMENT.RIGHT)
    rich(fp, "22AI741\t", 9, name="Calibri")
    rich(fp, "PAGE ", 9, color=RGBColor(0x44, 0x72, 0xC4), name="Calibri")
    field(fp, "PAGE", 9, RGBColor(0x44, 0x72, 0xC4))
    rich(fp, "\tDept. AI & ML", 9, name="Calibri")


def build(pages, toc=None, figs=None, tabs=None):
    R = Report(pages)
    front_matter(R)
    contents(R, toc or [], figs or [], tabs or [])
    main_section(R)
    body(R)
    return R


def pdf_pages(docx_path):
    out = docx_path.parent / "_aai_render"
    out.mkdir(exist_ok=True)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)],
                   check=True, capture_output=True)
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(out / (docx_path.stem + ".pdf")))
    texts = [pdf[i].get_textpage().get_text_range() for i in range(len(pdf))]
    pdf.close()
    return texts, out / (docx_path.stem + ".pdf")


def page_map(texts, keys):
    found = {}
    for t in texts:
        m = re.search(r"22AI741\s+PAGE\s+(\d+)", t)
        if not m:
            continue
        n = int(m.group(1))
        flat = re.sub(r"\s+", " ", t)
        for k in keys:
            if k not in found and re.sub(r"\s+", " ", k) in flat:
                found[k] = n
    return found


def prepare_annexures():
    from PIL import Image
    for name, pages, count in (("annex_wk.png", 5, 4), ("annex_po.png", 6, 6)):
        ims = [Image.open(A / f"annex_p{pages}_{k}.png").convert("RGB") for k in range(count)]
        w = max(i.width for i in ims)
        sheet = Image.new("RGB", (w, sum(i.height for i in ims)), "white")
        y = 0
        for im in ims:
            sheet.paste(im, (0, y))
            y += im.height
        sheet.save(A / name)


def main():
    prepare_annexures()
    r1 = build({})
    r1.doc.save(OUT_DOCX)
    keys = [k for _, _, k in r1.toc if k] + [k for _, _, k in r1.figs] + [k for _, _, k in r1.tabs]
    texts, _ = pdf_pages(OUT_DOCX)
    pages = page_map(texts, keys)
    missing = [k for k in keys if k not in pages]
    if missing:
        print("WARNING: page not found for", missing)
    toc = [(lvl if lvl >= 0 else 0, lab, key) for lvl, lab, key in r1.toc if lvl >= 0]
    r2 = build(pages, toc, r1.figs, r1.tabs)
    r2.doc.save(OUT_DOCX)
    texts, pdf = pdf_pages(OUT_DOCX)
    final = DOCS / "Case_Study_65_AAI_Report.pdf"
    final.write_bytes(pdf.read_bytes())
    print(OUT_DOCX)
    print(final, len(texts), "pages")


if __name__ == "__main__":
    main()
