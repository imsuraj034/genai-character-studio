"""Builds the case-study report (PDF) and presentation (PPTX) with live evaluation numbers.

Run:  python docs/build_docs.py
"""
import sys
from pathlib import Path

DOCS = Path(__file__).resolve().parent
sys.path.insert(0, str(DOCS.parent))
sys.path.insert(0, str(DOCS))

from fpdf import FPDF
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from evaluate import evaluate
from studio.consistency import RULES

STUDENT = "Suraj T"
USN = "1DS23AI058"
TITLE = "Case Study 65 - GenAI for Video Game Character Design"
SUBTITLE = "An AI-assisted game-design workflow with consistency checking, guardrails and human creative control"

NAVY, BLUE, GOLD, RED, GREEN, GREY = (32, 44, 74), (91, 110, 225), (212, 167, 44), (217, 83, 79), (46, 139, 87), (90, 90, 90)

PIPELINE = [
    ("Designer brief", NAVY), ("Input guardrail", RED), ("Cache, budget, routing", GREEN),
    ("Generator (Game Bible)", BLUE), ("Output guardrail", RED), ("Consistency check", BLUE),
    ("Auto-repair", BLUE), ("Human review", GREEN), ("Export + provenance", NAVY),
]

REQUIREMENTS = [
    ("Character concepts", "Name, appearance, palette, abilities, signature item, backstory, motivation, flaw, "
                           "catchphrase, a concept-art portrait drawn from the validated data (so the art matches "
                           "the Game Bible) and a ready-to-use prompt for an image model."),
    ("Dialogue & story variations", "Situation-aware dialogue in the character's personality voice; five story-arc "
                                    "types (Redemption, Betrayal, Tragic Sacrifice, Rise to Power, Hidden Heir)."),
    ("Consistency with rules & narrative", "A machine-readable Game Bible grounds every prompt; seven automated rules "
                                           "score every output; violated fields are regenerated automatically."),
    ("Copyright", "Fuzzy name matching against protected characters, franchise-term detection, signature-trait "
                  "combinations and 'in the style of <artist/studio>' detection."),
    ("Inappropriate content", "Age-rating content filter on both designer input and model output, plus "
                              "prompt-injection blocking; flagged text can be redacted."),
    ("Computational cost", "Token metering, response caching, task-based model routing, targeted (field-level) "
                           "repair, a session budget cap and a $0 offline/local engine."),
    ("Human creative control", "Nothing is exported without explicit designer approval; designers edit any field, "
                               "choose between 3 concept-art variations or upload their own artwork; every AI and "
                               "human version is stored with author, score and a diff."),
]


def latin(s):
    return (str(s).replace("—", "-").replace("–", "-").replace("→", "->")
            .replace("’", "'").replace("“", '"').replace("”", '"'))


# ======================================================================= PDF report
class Report(FPDF):
    def header(self):
        if self.page_no() > 1:
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(*GREY)
            self.cell(0, 6, latin(TITLE), align="R", new_x="LMARGIN", new_y="NEXT")
            self.ln(2)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(*GREY)
        self.cell(0, 6, f"Page {self.page_no()}", align="C")

    def h1(self, t):
        self.ln(3)
        self.set_font("Helvetica", "B", 15)
        self.set_text_color(*NAVY)
        self.cell(0, 9, latin(t), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*GOLD)
        self.set_line_width(0.6)
        self.line(self.l_margin, self.get_y(), self.l_margin + 30, self.get_y())
        self.ln(3)

    def h2(self, t):
        self.ln(1)
        self.set_font("Helvetica", "B", 11.5)
        self.set_text_color(*NAVY)
        self.cell(0, 7, latin(t), new_x="LMARGIN", new_y="NEXT")

    def para(self, t):
        self.set_font("Helvetica", "", 10.5)
        self.set_text_color(25, 25, 25)
        self.multi_cell(0, 5.6, latin(t), new_x="LMARGIN", new_y="NEXT")
        self.ln(1.5)

    def bullets(self, items):
        self.set_font("Helvetica", "", 10.5)
        self.set_text_color(25, 25, 25)
        for it in items:
            self.set_x(self.l_margin + 3)
            self.multi_cell(0, 5.6, latin(f"- {it}"), new_x="LMARGIN", new_y="NEXT")
        self.ln(1.5)

    def table(self, header, rows, widths):
        self.set_font("Helvetica", "B", 9)
        self.set_fill_color(*NAVY)
        self.set_text_color(255, 255, 255)
        for h, w in zip(header, widths):
            self.cell(w, 7, latin(h), border=0, fill=True, align="C")
        self.ln()
        self.set_font("Helvetica", "", 9)
        self.set_text_color(25, 25, 25)
        for i, r in enumerate(rows):
            self.set_fill_color(*((240, 243, 252) if i % 2 == 0 else (255, 255, 255)))
            line_h = 5
            # compute row height from the tallest cell
            heights = [len(self.multi_cell(w, line_h, latin(c), dry_run=True, output="LINES")) * line_h
                       for c, w in zip(r, widths)]
            h = max(heights) + 1
            if self.get_y() + h > self.page_break_trigger:
                self.add_page()
            x, y = self.get_x(), self.get_y()
            for c, w in zip(r, widths):
                self.rect(x, y, w, h, "F")
                self.set_xy(x, y + 0.5)
                self.multi_cell(w, line_h, latin(c))
                x += w
            self.set_xy(self.l_margin, y + h)
        self.ln(3)

    def pipeline_diagram(self):
        w, h, gap = 19.6, 15, 1.4
        x0, y = self.l_margin, self.get_y() + 2
        for i, (label, col) in enumerate(PIPELINE):
            x = x0 + i * (w + gap)
            self.set_fill_color(*col)
            self.rect(x, y, w, h, "F", round_corners=True, corner_radius=2)
            self.set_text_color(255, 255, 255)
            self.set_font("Helvetica", "B", 6.6)
            self.set_xy(x + 0.5, y + 2.5)
            self.multi_cell(w - 1, 3.4, latin(label), align="C")
        self.set_text_color(*GREY)
        self.set_font("Helvetica", "I", 8)
        self.set_xy(x0, y + h + 2)
        self.multi_cell(0, 4, "Game Bible feeds the Generator and the Consistency checker. If the checker finds "
                              "errors, Auto-repair regenerates only the broken fields and loops back through the "
                              "guardrails (max 2-3 rounds).", align="C")
        self.ln(3)


def build_report(rows, guard):
    pdf = Report()
    pdf.set_auto_page_break(True, 16)
    pdf.set_margins(16, 14, 16)

    # cover
    pdf.add_page()
    pdf.set_fill_color(*NAVY)
    pdf.rect(0, 0, 210, 297, "F")
    pdf.set_fill_color(*GOLD)
    pdf.rect(16, 92, 40, 1.4, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 28)
    pdf.set_xy(16, 100)
    pdf.multi_cell(178, 12, "GenAI Character\nDesign Studio")
    pdf.set_font("Helvetica", "", 13)
    pdf.set_xy(16, 132)
    pdf.multi_cell(178, 7, latin(TITLE + "\n\n" + SUBTITLE))
    pdf.set_font("Helvetica", "", 12)
    pdf.set_xy(16, 230)
    pdf.multi_cell(178, 7, f"Name: {STUDENT}\nUSN: {USN}")

    pdf.add_page()
    pdf.h1("1. Abstract")
    pdf.para("Game studios want to use generative AI to speed up character design, but raw model output is "
             "unreliable: it breaks the game's rules and lore, can imitate copyrighted characters, can produce "
             "content unsuitable for the target age rating, costs money per call, and risks taking creative "
             "control away from designers. This project builds a working AI-assisted design workflow that "
             "generates character concepts, dialogue and story variations, and wraps the generator in "
             "guardrails: a machine-readable Game Bible, input and output safety filters, a copyright "
             "detector, a rule-based consistency checker with automatic repair, cost controls, and a "
             "mandatory human review step with full version history. The system runs fully offline with no "
             "API key, and can switch to a free local LLM (Ollama) without code changes.")

    pdf.h1("2. Problem statement")
    pdf.para("A game-development studio wants to use GenAI to generate character concepts, dialogue and story "
             "variations. The problem is to develop an AI-assisted game-design workflow that generates "
             "content while maintaining consistency with game rules and narrative constraints, while also "
             "considering copyright, inappropriate content, computational cost and human creative control.")
    pdf.h2("Objectives")
    pdf.bullets([
        "Generate complete, usable character packages (concept, abilities, backstory, art prompt, dialogue, arcs).",
        "Ground all generation in an editable Game Bible and measure consistency automatically.",
        "Detect and fix rule violations without regenerating everything.",
        "Block unsafe or manipulative input and output; flag likely copyright problems.",
        "Make cost visible and controllable.",
        "Keep the human designer as the final decision-maker, with an audit trail.",
    ])

    pdf.h1("3. System architecture")
    pdf.pipeline_diagram()
    pdf.table(["Module", "File", "Responsibility"], [
        ["Game Bible", "studio/game_bible.py", "World rules, lore, factions, tech level, banned terms (JSON, editable in UI)"],
        ["Generator", "studio/engines.py", "Offline procedural engine or local LLM (Ollama); same interface"],
        ["Prompts", "studio/prompts.py", "Bible-grounded prompts for character, repair, dialogue, arcs"],
        ["Guardrails", "studio/safety.py", "Content filter, prompt-injection filter, copyright detector"],
        ["Consistency", "studio/consistency.py", "Rules R1-R7, score out of 100, error/warning severity"],
        ["Pipeline", "studio/pipeline.py", "Orchestrates all stages and the auto-repair loop"],
        ["Cost", "studio/cost.py", "Token metering, caching, model routing, budget cap"],
        ["Review", "studio/review.py", "Versioned human review: approve, edit, reject, diff"],
        ["Export", "studio/export.py", "JSON for game engines, PDF character sheet with provenance"],
        ["UI", "app.py", "Streamlit web app with 7 tabs"],
    ], [32, 42, 104])

    pdf.h1("4. How each requirement is met")
    pdf.table(["Requirement", "Design decision"], [[a, b] for a, b in REQUIREMENTS], [45, 133])

    pdf.h2("4.1 Consistency rules")
    pdf.table(["ID", "Rule", "Severity"], [
        [k, v, "warning" if k in ("R4", "R5", "R6") else "error"] for k, v in RULES.items()
    ], [14, 134, 30])
    pdf.para("Score = 100 - 20 x errors - 7 x warnings. Errors block approval unless the designer explicitly "
             "overrides and takes responsibility, which is recorded in the version history.")

    pdf.h2("4.2 Why an offline engine?")
    pdf.para("Real LLMs make characteristic mistakes: they mix up magic systems, add anachronistic technology, "
             "contradict lore, reuse famous names and copy artists' styles. The offline engine reproduces these "
             "mistakes on purpose. The higher the creativity (temperature) setting, the more often it does so. "
             "This makes the guardrails demonstrable and testable without an API key. The same pipeline runs "
             "unchanged on a local open-source LLM through Ollama.")

    pdf.h1("5. Evaluation")
    pdf.para(f"Each temperature setting was run on {rows[0]['runs']} random briefs (random faction, class, role "
             "and personality). 'Draft' is the raw generator output; 'final' is after the automatic repair loop.")
    pdf.table(["Temp.", "Drafts with violations", "Avg draft score", "Avg final score", "Finals with errors",
               "Avg repair rounds"],
              [[r["temperature"], f"{r['drafts_with_violations_pct']}%", r["avg_draft_score"], r["avg_final_score"],
                f"{r['final_with_errors_pct']}%", r["avg_repair_rounds"]] for r in rows],
              [18, 36, 30, 30, 32, 32])
    hi = rows[-1]
    pdf.para(f"At the highest creativity, {hi['drafts_with_violations_pct']}% of raw drafts broke at least one rule "
             f"(average score {hi['avg_draft_score']}), yet after auto-repair the average score was "
             f"{hi['avg_final_score']} with {hi['final_with_errors_pct']}% of characters still containing errors. "
             "Higher creativity costs more because more repair calls are needed. That is a real trade-off for the "
             "studio to tune.")
    pdf.h2("Guardrail accuracy (test sets)")
    pdf.table(["Test", "Result"], [
        ["Unsafe / injection briefs blocked", guard["unsafe_blocked"]],
        ["Safe briefs wrongly blocked", guard["safe_false_blocks"]],
        ["Copycat names flagged (e.g. Pikachoo, Kratoz)", guard["copycat_names_flagged"]],
        ["Original generated names wrongly flagged", guard["original_names_false_flags"]],
    ], [120, 58])
    pdf.para("The one borderline flag ('Maro' is 89% similar to 'Mario') shows why copyright results are flags for "
             "a human to judge rather than automatic blocks.")
    pdf.h2("Cost")
    pdf.para(f"With task-based routing (small model for repair and dialogue, medium model for creative work) the "
             f"estimated cost was ${hi['cost_per_char_usd']:.4f} per character versus "
             f"${hi['large_model_cost_per_char_usd']:.4f} if the largest model were used for everything, about "
             f"{hi['large_model_cost_per_char_usd'] / hi['cost_per_char_usd']:.0f}x cheaper. Repeated requests are "
             "served from cache at $0, and the offline/local engine costs $0. Prices are illustrative tiers "
             "configurable in the app.")
    pdf.h2("Automated tests")
    pdf.para("8 automated tests (pytest) cover every faction/class/personality combination, drift-and-repair, "
             "input blocking, copyright detection, caching, the budget cap, dialogue/story checks, the review "
             "workflow, export, and a headless run of the full Streamlit UI.")

    pdf.h1("6. Ethics and responsible use")
    pdf.bullets([
        "Human authorship: AI output is a draft; designers approve, edit or reject, and every change is attributed.",
        "IP respect: the system discourages imitation of named artists or studios and flags look-alike characters.",
        "Age-appropriate content enforced at input and output; prompt-injection attempts are blocked.",
        "Transparency: provenance (AI vs human, version, score) is exported with every character.",
        "Sustainability: caching, routing and targeted repair reduce compute use.",
    ])

    pdf.h1("7. Limitations and future work")
    pdf.bullets([
        "Keyword and pattern filters can miss paraphrases. A trained moderation model and embedding-based "
        "similarity search would improve recall.",
        "Copyright checks use a sample list of protected names; a production system would use a licensed database.",
        "Integrate an image model (e.g. a local Stable Diffusion) to turn the concept-art prompt into images, with "
        "the same style guardrails.",
        "Multi-user review with roles (writer, art director, legal) and comment threads.",
        "Fine-tune a small model on the studio's own approved characters for a consistent house voice.",
    ])

    pdf.h1("8. Conclusion")
    pdf.para("The project shows that GenAI can be used productively in character design when it is surrounded by "
             "the right engineering: grounding in a Game Bible, automated verification and repair, safety and "
             "copyright guardrails, cost controls, and a human who always has the final say. The result is "
             "faster ideation without giving up quality, safety or creative ownership.")

    pdf.h1("9. How to run")
    pdf.para("1) Install Python 3.10+.  2) Double-click run.bat (or: pip install -r requirements.txt, then "
             "streamlit run app.py).  3) Optional: install Ollama and run 'ollama pull llama3.2' to use a local LLM.  "
             "4) Tests: python -m pytest -q")

    pdf.h1("References")
    pdf.bullets([
        "PEGI - Pan European Game Information age rating system, pegi.info.",
        "Park, J. S. et al. (2023). Generative Agents: Interactive Simulacra of Human Behavior. UIST '23.",
        "Streamlit documentation, docs.streamlit.io.",
        "Ollama - run open-source large language models locally, ollama.com.",
    ])
    out = DOCS / "Case_Study_65_Report.pdf"
    pdf.output(str(out))
    return out


# ======================================================================= PPTX slides
def rgb(c):
    return RGBColor(*c)


def build_slides(rows, guard):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def bg(slide, color=(250, 250, 252)):
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = rgb(color)

    def text(slide, x, y, w, h, s, size=18, bold=False, color=(30, 30, 30), align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = align
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = rgb(color)
            p.space_after = Pt(6)
        return tb

    def box(slide, x, y, w, h, fill, label, size=13, font=(255, 255, 255)):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        s.line.fill.background()
        tf = s.text_frame
        tf.word_wrap = True
        tf.text = label
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(size)
            p.font.bold = True
            p.font.color.rgb = rgb(font)
        return s

    def titled(title, kicker=None):
        s = prs.slides.add_slide(blank)
        bg(s)
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)).fill.solid()
        s.shapes[-1].fill.fore_color.rgb = rgb(GOLD)
        s.shapes[-1].line.fill.background()
        if kicker:
            text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), 12, True, BLUE)
        text(s, 0.6, 0.65, 12, 0.9, title, 30, True, NAVY)
        return s

    def bullet_slide(title, items, kicker=None, size=20):
        s = titled(title, kicker)
        text(s, 0.8, 1.8, 11.8, 5.2, "\n".join(f"•  {i}" for i in items), size)
        return s

    # 1 title
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    text(s, 0.8, 2.0, 11.5, 1.2, "GenAI Character Design Studio", 46, True, (255, 255, 255))
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(3.25), Inches(1.6), Inches(0.07)).fill.solid()
    s.shapes[-1].fill.fore_color.rgb = rgb(GOLD)
    s.shapes[-1].line.fill.background()
    text(s, 0.8, 3.5, 11.5, 1.2, TITLE + "\n" + SUBTITLE, 18, False, (210, 220, 245))
    text(s, 0.8, 5.9, 11.5, 0.9, f"{STUDENT}   |   USN: {USN}", 16, False, (255, 255, 255))

    # 2 problem
    bullet_slide("The problem", [
        "Studios want GenAI to speed up character concepts, dialogue and story variations",
        "But raw AI output breaks game rules and lore (wrong magic, guns in a medieval world, living dragons...)",
        "It can copy famous characters or artists: copyright risk",
        "It can produce content unsuitable for the age rating",
        "Every call costs money and compute",
        "Designers must keep creative control: AI should assist, not replace",
    ], "Case Study 65")

    # 3 solution pipeline
    s = titled("Our solution: a guarded design pipeline", "Architecture")
    w, gap, y = 1.28, 0.1, 2.6  # 9 boxes fit in 12.4" of a 13.33" slide
    for i, (label, col) in enumerate(PIPELINE):
        box(s, 0.45 + i * (w + gap), y, w, 1.25, col, label, 12)
    box(s, 4.6, 4.5, 4.2, 0.9, GOLD, "📖 Game Bible: rules · lore · factions · banned content", 13, NAVY)
    text(s, 0.6, 5.7, 12, 1.2, "The Game Bible grounds both generation and checking. Errors trigger auto-repair of only the "
                               "broken fields, which then pass through the guardrails again.", 16, False, GREY)

    # 4 game bible
    bullet_slide("Game Bible = single source of truth", [
        "World: Aethervale, a high-fantasy RPG rated PEGI 12, set at a medieval tech level",
        "4 factions, each with allowed magic elements, colours, naming rules and a rival",
        "Lore facts, e.g. 'Dragons are extinct'; contradictions are detected with patterns",
        "Forbidden tech list: gun, laser, robot, electricity...",
        "Editable in the app: change a rule and every new generation and check follows it immediately",
        "Injected into every prompt, so the model and the checker see the same rules",
    ], "Consistency")

    # 5 consistency + repair
    s = titled("Consistency checking and auto-repair", "Consistency")
    text(s, 0.6, 1.7, 6.2, 5, "\n".join(f"{k}  {v}" for k, v in RULES.items()), 15)
    hi = rows[-1]
    box(s, 7.3, 1.9, 2.6, 1.6, RED, f"Raw AI draft\n{hi['avg_draft_score']}/100", 20)
    box(s, 10.2, 1.9, 2.6, 1.6, GREEN, f"After auto-repair\n{hi['avg_final_score']}/100", 20)
    text(s, 7.3, 3.8, 5.5, 3, f"At max creativity, {hi['drafts_with_violations_pct']}% of drafts broke a rule.\n"
                             f"After repair, {hi['final_with_errors_pct']}% still had errors.\n"
                             "Only broken fields are regenerated, which keeps cost low and keeps the parts the designer liked.",
         15, False, GREY)

    # 6 guardrails
    s = titled("Guardrails: safety and copyright", "Responsible AI")
    text(s, 0.6, 1.7, 6, 5, "Inappropriate content\n• Checked on input and output\n• Categories: violence, sexual, "
                            "self-harm, drugs, hate, profanity\n• Prompt-injection attempts blocked\n\nCopyright\n"
                            "• Fuzzy name match vs protected characters\n• Franchise terms (Jedi, Hogwarts...)\n"
                            "• Signature-trait combos (red cap + plumber)\n• 'In the style of <artist>' detection", 16)
    for i, (k, v) in enumerate([("Unsafe briefs blocked", guard["unsafe_blocked"]),
                                ("Safe briefs wrongly blocked", guard["safe_false_blocks"]),
                                ("Copycat names flagged", guard["copycat_names_flagged"]),
                                ("Original names false-flagged", guard["original_names_false_flags"])]):
        box(s, 7.2 + (i % 2) * 2.95, 1.9 + (i // 2) * 2.0, 2.75, 1.7, NAVY if i % 2 == 0 else BLUE, f"{v}\n{k}", 16)

    # 7 cost
    s = titled("Controlling computational cost", "Cost")
    text(s, 0.6, 1.7, 6.5, 5, "• Token metering on every call\n• Model routing: small model for repair and dialogue, "
                              "medium for creative work\n• Response cache: repeat requests cost $0\n• Field-level repair "
                              "instead of full regeneration\n• Session budget cap\n• Offline / local engine: $0 for prototyping", 17)
    box(s, 7.6, 2.0, 2.4, 1.8, GREEN, f"${hi['cost_per_char_usd']:.4f}\nper character (routed)", 15)
    box(s, 10.3, 2.0, 2.4, 1.8, RED, f"${hi['large_model_cost_per_char_usd']:.4f}\nlargest model only", 15)
    text(s, 7.6, 4.1, 5.2, 1.5, f"≈ {hi['large_model_cost_per_char_usd'] / hi['cost_per_char_usd']:.0f}x cheaper. "
                                "Prices are illustrative tiers and editable in the app.", 15, False, GREY)

    # 8 human control
    bullet_slide("Human creative control", [
        "AI output is always a draft: status 'pending' until a designer acts",
        "Approve / Reject / Regenerate / Edit any field; edits are re-checked automatically",
        "Concept art: pick 1 of 3 AI variations, ask for more, or upload your own artwork",
        "Approving with rule errors requires an explicit, logged override",
        "Full version history: author (AI or Designer), action, score, note, and a diff between versions",
        "Only approved characters can be exported, with provenance included",
    ], "Human in the loop")

    # 9 demo
    bullet_slide("Live demo", [
        "1. Normal brief: a clean character that passes every check",
        "2. High creativity: the AI breaks rules, then auto-repair takes the score from ~20 to 100",
        "3. Unsafe brief and prompt injection: blocked before reaching the model",
        "4. Copycat name 'Pikachoo': copyright flag",
        "5. Dialogue and 3 story variations for the same character",
        "6. Designer picks concept art, edits, approves, then exports a PDF character sheet with the art",
    ], "Demo")

    # 10 results
    s = titled("Results", "Evaluation")
    hdr = ["Temp", "Drafts w/ violations", "Draft score", "Final score", "Finals w/ errors"]
    tbl = s.shapes.add_table(len(rows) + 1, len(hdr), Inches(0.8), Inches(1.9), Inches(11.7), Inches(3)).table
    for j, h in enumerate(hdr):
        tbl.cell(0, j).text = h
    for i, r in enumerate(rows, 1):
        vals = [r["temperature"], f"{r['drafts_with_violations_pct']}%", r["avg_draft_score"],
                r["avg_final_score"], f"{r['final_with_errors_pct']}%"]
        for j, v in enumerate(vals):
            tbl.cell(i, j).text = str(v)
    for row in tbl.rows:
        for c in row.cells:
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(16)
    text(s, 0.8, 5.4, 11.7, 1.2, "8 automated tests passing, including a headless run of the whole UI.", 16, False, GREY)

    # 11 future
    bullet_slide("Limitations and future work", [
        "Keyword filters can miss paraphrases: add a trained moderation model and embedding similarity",
        "Use a licensed IP database instead of a sample list of protected names",
        "Generate concept art with a local image model, using the same style guardrails",
        "Multi-role review (writer, art director, legal) with comments",
        "Fine-tune on the studio's approved characters for a consistent house style",
    ], "Next steps")

    # 12 thanks
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    text(s, 0.8, 2.7, 11.5, 1.2, "Thank you", 48, True, (255, 255, 255))
    text(s, 0.8, 3.9, 11.5, 1, "Questions?", 22, False, (210, 220, 245))

    out = DOCS / "Case_Study_65_Slides.pptx"
    prs.save(str(out))
    return out


if __name__ == "__main__":
    rows, guard = evaluate()
    print(build_report(rows, guard))
    print(build_slides(rows, guard))
